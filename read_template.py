from openpyxl import load_workbook
import json

# Load the template file
wb = load_workbook('templete/client_import_template_validated (1).xlsx')
print('Sheet names:', wb.sheetnames)

# Get the Client Data sheet
ws = wb['Client Data'] if 'Client Data' in wb.sheetnames else wb[wb.sheetnames[0]]
print('\nActive sheet:', ws.title)

# Get headers
headers = []
for cell in ws[1]:
    if cell.value:
        headers.append(cell.value)

print('\nHeaders:')
print(json.dumps(headers, indent=2))

# Get first few data rows
print('\nFirst 3 data rows:')
for i in range(2, min(5, ws.max_row + 1)):
    row_data = {}
    for idx, cell in enumerate(ws[i]):
        if idx < len(headers):
            row_data[headers[idx]] = cell.value
    print(json.dumps(row_data, indent=2, default=str))
    print('---')

print(f'\nTotal rows: {ws.max_row}')
print(f'Total columns: {len(headers)}')
